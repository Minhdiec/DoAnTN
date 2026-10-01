from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Bao_cao_chuyen_de_TN_Nguyen_Thi_Hong_Minh_25410256 - Copy.docx"
OUT = ROOT / "Bao_cao_chuyen_de_TN_Nguyen_Thi_Hong_Minh_25410256 - Copy.final.docx"


FIGURE_LIST = [
    "Hình 3.1. Sơ đồ use case tổng quát của hệ thống",
    "Hình 3.2. Sơ đồ ERD cơ sở dữ liệu",
    "Hình 5.1. Giao diện đăng ký tài khoản",
    "Hình 5.2. Phần đầu trang chủ",
    "Hình 5.3. Khối Khám phá sản phẩm trên trang chủ",
    "Hình 5.4. Trang chi tiết sản phẩm Dublin",
    "Hình 5.5. Đánh giá sau khi gửi, kèm nhãn cảm xúc",
    "Hình 5.6. Cửa sổ thử kính ảo với mẫu Dublin",
    "Hình 5.7. Giỏ hàng của thành viên",
    "Hình 5.8. Cửa sổ chọn hình thức thanh toán và đơn vị vận chuyển",
    "Hình 5.9. Trang Đơn hàng của tôi sau khi thanh toán",
    "Hình 5.10. Danh sách đánh giá kèm thống kê cảm xúc",
]

TABLE_LIST = [
    "Bảng 3.1. Bảng đặc tả use case",
    "Bảng 3.2. Danh sách các bảng dữ liệu",
    "Bảng 3.3. Cấu trúc bảng User (accounts)",
    "Bảng 3.4. Cấu trúc bảng Category (products)",
    "Bảng 3.5. Cấu trúc bảng Product (products)",
    "Bảng 3.6. Cấu trúc bảng ProductImage (products)",
    "Bảng 3.7. Cấu trúc bảng Cart (cart)",
    "Bảng 3.8. Cấu trúc bảng CartItem (cart)",
    "Bảng 3.9. Cấu trúc bảng Order (orders)",
    "Bảng 3.10. Cấu trúc bảng OrderItem (orders)",
    "Bảng 3.11. Cấu trúc bảng Favorite (favorites)",
    "Bảng 3.12. Cấu trúc bảng Favorite_Product (favorites)",
    "Bảng 3.13. Cấu trúc bảng Review (reviews)",
    "Bảng 3.14. Cấu trúc bảng ReviewMedia (reviews)",
    "Bảng 3.15. Cấu trúc bảng Wallet (wallet)",
    "Bảng 3.16. Cấu trúc bảng Transaction (wallet)",
    "Bảng 3.17. Cấu trúc bảng GlassesOverlay (tryon)",
    "Bảng 4.1. Các thư viện chính và vai trò trong dự án",
    "Bảng 4.2. URL và router của toàn hệ thống",
    "Bảng 4.3. Giải thích các hàm và phương thức chính",
    "Bảng 4.4. Luồng xử lý một khung hình giữa trình duyệt và server",
    "Bảng 4.5. FPS trước và sau khi áp dụng tối ưu",
    "Bảng 4.6. So sánh đóng góp của từng kỹ thuật tối ưu",
    "Bảng 4.7. So sánh 3 mô hình phân loại cảm xúc",
    "Bảng 5.1. Các mục quản lý trong trang quản trị Django Admin",
]

REFERENCES = [
    "[1] Django Software Foundation. (2026). Django documentation, version 4.2. https://docs.djangoproject.com/en/4.2/",
    "[2] Django Software Foundation. (2026). The Django admin site. https://docs.djangoproject.com/en/4.2/ref/contrib/admin/",
    "[3] Django Channels Project. (2026). Django Channels documentation. https://channels.readthedocs.io/",
    "[4] Google. (2026). MediaPipe Face Landmarker task guide. https://developers.google.com/mediapipe/solutions/vision/face_landmarker",
    "[5] OpenCV Team. (2026). OpenCV-Python tutorials and documentation. https://docs.opencv.org/",
    "[6] Casiez, G., Roussel, N., & Vogel, D. (2012). 1 Euro Filter: A simple speed-based low-pass filter for noisy input in interactive systems. CHI 2012. https://cristal.univ-lille.fr/~casiez/1euro/",
    "[7] Scikit-learn Developers. (2026). scikit-learn user guide: Feature extraction and linear models. https://scikit-learn.org/stable/",
    "[8] underthesea Contributors. (2026). underthesea: Vietnamese NLP toolkit documentation. https://underthesea.readthedocs.io/",
    "[9] PyMySQL Contributors. (2026). PyMySQL documentation. https://pymysql.readthedocs.io/",
    "[10] Oracle. (2026). MySQL 8.0 reference manual. https://dev.mysql.com/doc/refman/8.0/en/",
]

REMOVE_FIGURES = {
    "5.2",
    "5.3",
    "5.4",
    "5.7",
    "5.9",
    "5.12",
    "5.13",
    "5.17",
    "5.18",
    "5.19",
    "5.20",
    "5.21",
    "5.22",
    "5.23",
    "5.24",
    "5.25",
    "5.26",
    "5.27",
    "5.28",
    "5.29",
}

FIGURE_RENUMBER = {
    "5.1": "5.1",
    "5.5": "5.2",
    "5.6": "5.3",
    "5.8": "5.4",
    "5.10": "5.5",
    "5.11": "5.6",
    "5.14": "5.7",
    "5.15": "5.8",
    "5.16": "5.9",
    "5.30": "5.10",
}


def block_items(doc: DocumentType):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def text_of(block) -> str:
    if isinstance(block, Paragraph):
        return " ".join(block.text.split())
    if isinstance(block, Table):
        return " ".join(cell.text for row in block.rows for cell in row.cells).strip()
    return ""


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def has_drawing(paragraph: Paragraph) -> bool:
    xml = paragraph._p.xml
    return "<w:drawing" in xml or "<w:pict" in xml


def insert_paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    new_para.add_run(text)
    return new_para


def replace_text(paragraph: Paragraph, replacements: dict[str, str]) -> None:
    if not paragraph.runs:
        return
    full = "".join(run.text for run in paragraph.runs)
    changed = full
    for old, new in replacements.items():
        changed = changed.replace(old, new)
    if changed == full:
        return
    paragraph.runs[0].text = changed
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_section(doc: DocumentType, start_prefix: str, stop_prefix: str) -> None:
    deleting = False
    to_delete = []
    for block in list(block_items(doc)):
        text = text_of(block)
        if isinstance(block, Paragraph) and text.startswith(start_prefix):
            deleting = True
        if deleting and isinstance(block, Paragraph) and text.startswith(stop_prefix):
            deleting = False
        if deleting:
            to_delete.append(block._element)
    for element in to_delete:
        remove_element(element)


def replace_list_between_headings(doc: DocumentType, heading: str, next_heading: str, entries: list[str]) -> None:
    blocks = list(block_items(doc))
    start = stop = None
    for i, block in enumerate(blocks):
        txt = text_of(block)
        if isinstance(block, Paragraph) and txt == heading:
            start = i
        elif start is not None and isinstance(block, Paragraph) and txt == next_heading:
            stop = i
            break
    if start is None or stop is None:
        return
    heading_para = blocks[start]
    for block in blocks[start + 1 : stop]:
        remove_element(block._element)
    cursor = heading_para
    for entry in entries:
        cursor = insert_paragraph_after(cursor, entry)


def replace_references(doc: DocumentType) -> None:
    blocks = list(block_items(doc))
    start = None
    for i, block in enumerate(blocks):
        if isinstance(block, Paragraph) and text_of(block) == "TÀI LIỆU THAM KHẢO":
            start = i
            break
    if start is None:
        return
    heading = blocks[start]
    for block in blocks[start + 1 :]:
        remove_element(block._element)
    cursor = heading
    for entry in REFERENCES:
        cursor = insert_paragraph_after(cursor, entry, "List Paragraph")


def remove_selected_figures(doc: DocumentType) -> None:
    blocks = list(block_items(doc))
    to_delete = set()
    in_chapter_5 = False
    for i, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        txt = text_of(block)
        if txt == "CHƯƠNG 5: GIAO DIỆN VÀ DEMO CHỨC NĂNG":
            in_chapter_5 = True
            continue
        if in_chapter_5 and txt == "CHƯƠNG 6: KẾT LUẬN":
            in_chapter_5 = False
        if not in_chapter_5:
            continue
        m = re.match(r"Hình (5\.\d+)\.", txt)
        if not m or m.group(1) not in REMOVE_FIGURES:
            continue
        to_delete.add(block._element)

        # Remove the image paragraph immediately before the caption.
        j = i - 1
        while j >= 0 and isinstance(blocks[j], Paragraph) and not text_of(blocks[j]) and not has_drawing(blocks[j]):
            to_delete.add(blocks[j]._element)
            j -= 1
        if j >= 0 and isinstance(blocks[j], Paragraph) and has_drawing(blocks[j]):
            to_delete.add(blocks[j]._element)

        # Remove the short explanatory paragraph after the deleted caption.
        k = i + 1
        while k < len(blocks):
            nxt = blocks[k]
            if not isinstance(nxt, Paragraph):
                break
            nxt_text = text_of(nxt)
            if not nxt_text:
                to_delete.add(nxt._element)
                k += 1
                continue
            style_name = nxt.style.name if nxt.style is not None else ""
            if style_name.startswith("Heading") or re.match(r"Hình \d+\.\d+\.", nxt_text):
                break
            to_delete.add(nxt._element)
            break

    for element in list(to_delete):
        remove_element(element)


def update_caption_numbers(doc: DocumentType) -> None:
    replacements = {
        "ở Hình 5.11 và 5.12": "ở Hình 5.6",
        "xem chi tiết bên dưới": "xem Hình 5.10",
        "Ảnh được chụp khi chạy ứng dụng trên máy cục bộ với dữ liệu mẫu (2 danh mục, 4 sản phẩm, 60 đánh giá do lệnh seed_reviews tạo). Các thao tác đăng ký, đặt hàng, huỷ đơn và đánh giá trong ảnh được thực hiện trực tiếp trên giao diện.": (
            "Chương này chỉ giữ các màn hình đại diện cho luồng chính: đăng ký, duyệt sản phẩm, thử kính ảo, giỏ hàng, đặt hàng và quản trị. "
            "Các thao tác trong ảnh được thực hiện trực tiếp trên ứng dụng chạy cục bộ với dữ liệu mẫu."
        ),
    }
    for block in block_items(doc):
        if isinstance(block, Paragraph):
            replace_text(block, replacements)
            txt = text_of(block)
            for old, new in sorted(FIGURE_RENUMBER.items(), key=lambda item: -len(item[0])):
                if f"Hình {old}." in txt or f"Hình {old} " in txt:
                    replace_text(block, {f"Hình {old}.": f"Hình {new}.", f"Hình {old} ": f"Hình {new} "})


def apply_global_text_fixes(doc: DocumentType) -> None:
    fixes = {
        "Bấm chuột phải vào bảng dưới, chọn Update Field hoặc Cập nhật trường để tự động điền đúng số trang sau khi hoàn tất chỉnh sửa nội dung.": "",
        "Sơ đồ được sinh từ các model Django của 8 app, gồm 14 bảng.": "Sơ đồ được sinh từ cơ sở dữ liệu và model Django của 8 app, gồm 15 bảng nghiệp vụ.",
        "khong co": "không có",
        "không mã hoá": "không mã hóa",
        "mã hoá": "mã hóa",
        "chuẩn hoá": "chuẩn hóa",
        "từ khoá": "từ khóa",
        "tuỳ": "tùy",
        "toạ độ": "tọa độ",
        "lựa chọn": "lựa chọn",
        "Đổi sang mẫu Incantation": "Đổi mẫu kính trong cửa sổ thử kính ảo",
    }
    for block in block_items(doc):
        if isinstance(block, Paragraph):
            if text_of(block) == next(iter(fixes)):
                remove_element(block._element)
                continue
            replace_text(block, fixes)
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_text(paragraph, fixes)


def update_cover(doc: DocumentType) -> None:
    if not doc.tables:
        return
    cell = doc.tables[0].cell(0, 0)
    for paragraph in cell.paragraphs:
        replace_text(
            paragraph,
            {
                "[ TÊN TRƯỜNG, vui lòng điền ]": "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN",
                "[ KHOA CÔNG NGHỆ THÔNG TIN, vui lòng điền ]": "KHOA CÔNG NGHỆ THÔNG TIN",
                "Lớp: [ Điền lớp ]": "",
            },
        )


def renumber_chapter_6(doc: DocumentType) -> None:
    replacements = {
        "6.3. Đóng góp của đề tài": "6.2. Đóng góp của đề tài",
        "6.3.1. Tính đổi mới": "6.2.1. Tính đổi mới",
        "6.3.2. Tính thực tiễn": "6.2.2. Tính thực tiễn",
        "6.4. Hạn chế": "6.3. Hạn chế",
        "6.5. Hướng phát triển": "6.4. Hướng phát triển",
    }
    for block in block_items(doc):
        if isinstance(block, Paragraph):
            replace_text(block, replacements)


def resize_images(doc: DocumentType) -> None:
    for idx, shape in enumerate(doc.inline_shapes, 1):
        old_w, old_h = shape.width, shape.height
        ratio = old_h / old_w if old_w else 1
        if idx <= 2:
            target_w = Inches(5.8)
        else:
            target_w = Inches(4.65)
        target_h = int(target_w * ratio)
        max_h = Inches(5.6)
        if target_h > max_h:
            target_h = max_h
            target_w = int(target_h / ratio)
        shape.width = target_w
        shape.height = target_h


def tighten_paragraph_spacing(doc: DocumentType) -> None:
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            paragraph.paragraph_format.space_after = 0
        elif text_of(paragraph):
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.line_spacing = 1.0


def main() -> None:
    doc = Document(SRC)
    update_cover(doc)
    remove_selected_figures(doc)
    delete_section(doc, "6.2. Đối chiếu", "6.3. Đóng góp")
    renumber_chapter_6(doc)
    apply_global_text_fixes(doc)
    update_caption_numbers(doc)
    replace_list_between_headings(doc, "DANH MỤC HÌNH VẼ", "DANH MỤC BẢNG", FIGURE_LIST)
    replace_list_between_headings(doc, "DANH MỤC BẢNG", "DANH MỤC TỪ VIẾT TẮT", TABLE_LIST)
    replace_references(doc)
    resize_images(doc)
    tighten_paragraph_spacing(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
